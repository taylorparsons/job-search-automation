#!/usr/bin/env python3
# Reverse-Recruiter Job Search System - tailor
# Copyright (C) 2026 Jennifer McKinney
# This program is free software: you can redistribute it and/or modify it under
# the terms of the GNU General Public License v3.0 as published by the Free
# Software Foundation. Distributed WITHOUT ANY WARRANTY. See
# <https://www.gnu.org/licenses/>. Commercial licensing available from the author.
"""Generate tailored application materials from a local profile or career bank.

Usage:
    python3 tailor.py --config ../config/config.json --role role.example.json --out ./output

Inputs:
    config.json  - your profile (see config.example.json). Uses the "operator" block,
                   plus optional "summary" (str) and "resume" {skills[], experience[]}.
    role.json    - the target role: {"company","title","hiring_manager"?,"highlights"?[]}

Outputs (in --out):
    CoverLetter_<Company>.docx
    Resume_<YourName>.docx

Career-bank mode delegates evidence selection and rendering to a local
``career_cli.py`` supplied by the operator. It never copies a career bank into
this repository.
"""
__version__ = "1.0.0"

import argparse
import importlib.util
import json
import sys
from datetime import date
from pathlib import Path
from types import ModuleType
from typing import Any


def load_json(path: str) -> dict[str, Any]:
    p = Path(path)
    if not p.is_file():
        sys.exit(f"File not found: {path}")
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        sys.exit(f"Invalid JSON in {path}: {exc}")


def _docx_dependencies() -> tuple[Any, Any]:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise SystemExit("python-docx is required for DOCX output. Install it with: pip install python-docx") from exc
    return Document, Pt


def _para(doc: Any, text: str, size: int = 11, bold: bool = False, space_after: int = 8) -> None:
    _, Pt = _docx_dependencies()
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    para.paragraph_format.space_after = Pt(space_after)


def build_cover_letter(cfg: dict[str, Any], role: dict[str, Any], outdir: Path) -> Path:
    Document, _ = _docx_dependencies()
    op = cfg.get("operator", {})
    company = role.get("company", "the company")
    title = role.get("title", "the role")
    greeting_name = role.get("hiring_manager") or "Hiring Team"
    highlights = role.get("highlights") or []
    summary = cfg.get("summary", "").strip()

    doc = Document()
    _para(doc, date.today().strftime("%B %d, %Y"), space_after=12)
    _para(doc, f"Dear {greeting_name},", space_after=10)
    _para(doc, f"I'm writing to express my interest in the {title} role at {company}.", space_after=10)
    if summary:
        _para(doc, summary, space_after=10)
    if highlights:
        _para(doc, f"A few reasons I'd be a strong fit for {company}:", bold=True, space_after=6)
        for h in highlights:
            b = doc.add_paragraph(style="List Bullet")
            b.add_run(str(h)).font.size = Pt(11)
    _para(
        doc,
        f"I'd welcome the chance to discuss how I can contribute to {company}. "
        "Thank you for your time and consideration.",
        space_after=12,
    )
    _para(doc, "Sincerely,", space_after=2)
    _para(doc, op.get("name", "Your Name"), bold=True, space_after=2)
    contact = " | ".join(x for x in [op.get("email", ""), op.get("phone", ""), op.get("linkedin", "")] if x)
    if contact:
        _para(doc, contact, size=10)

    safe = "".join(ch for ch in company if ch.isalnum() or ch in " -_").strip().replace(" ", "_")
    out = outdir / f"CoverLetter_{safe or 'Company'}.docx"
    doc.save(out)
    return out


def build_resume(cfg: dict[str, Any], outdir: Path) -> Path:
    Document, _ = _docx_dependencies()
    op = cfg.get("operator", {})
    resume = cfg.get("resume", {}) or {}

    doc = Document()
    _para(doc, op.get("name", "Your Name"), size=18, bold=True, space_after=2)
    contact = " | ".join(x for x in [op.get("email", ""), op.get("phone", ""), op.get("linkedin", ""), op.get("location_line", "")] if x)
    if contact:
        _para(doc, contact, size=10, space_after=10)
    if cfg.get("summary"):
        _para(doc, "SUMMARY", bold=True, space_after=4)
        _para(doc, cfg["summary"], space_after=10)
    if resume.get("skills"):
        _para(doc, "SKILLS", bold=True, space_after=4)
        _para(doc, ", ".join(str(s) for s in resume["skills"]), space_after=10)
    if resume.get("experience"):
        _para(doc, "EXPERIENCE", bold=True, space_after=4)
        for job in resume["experience"]:
            line = f"{job.get('title','')} - {job.get('company','')} ({job.get('dates','')})".strip(" -")
            _para(doc, line, bold=True, space_after=2)
            for bullet in job.get("bullets", []):
                b = doc.add_paragraph(style="List Bullet")
                b.add_run(str(bullet)).font.size = Pt(11)
    else:
        _para(doc, "Add a \"resume\" block to config.json to populate experience.", size=10)

    name_safe = "".join(ch for ch in op.get("name", "Resume") if ch.isalnum()) or "Resume"
    out = outdir / f"Resume_{name_safe}.docx"
    doc.save(out)
    return out


def load_career_cli(path: str) -> ModuleType:
    module_path = Path(path).expanduser().resolve()
    if not module_path.is_file():
        raise SystemExit(f"Career CLI not found: {module_path}")
    spec = importlib.util.spec_from_file_location("local_career_cli", module_path)
    if spec is None or spec.loader is None:
        raise SystemExit(f"Could not load career CLI: {module_path}")
    module = importlib.util.module_from_spec(spec)
    # Dataclasses in the Career CLI resolve postponed annotations through
    # sys.modules during import, so register the dynamically loaded module
    # before executing it.
    sys.modules[spec.name] = module
    try:
        spec.loader.exec_module(module)
    except Exception:
        sys.modules.pop(spec.name, None)
        raise
    return module


def _load_job_description(role: dict[str, Any], role_path: Path) -> tuple[str, str]:
    inline = str(role.get("job_description") or "").strip()
    if inline:
        # Never persist a pasted job description beside the role or in the
        # repository. Job descriptions can be sensitive, and the text is
        # already passed directly to the local Career CLI.
        return inline, f"{role_path.name} (inline job_description)"
    raw_path = str(role.get("job_description_path") or "").strip()
    if not raw_path:
        raise SystemExit("Career mode requires role.job_description or role.job_description_path")
    jd_path = Path(raw_path).expanduser()
    if not jd_path.is_absolute():
        jd_path = role_path.parent / jd_path
    try:
        return jd_path.read_text(encoding="utf-8"), str(jd_path)
    except FileNotFoundError as exc:
        raise SystemExit(f"Job description not found: {jd_path}") from exc


def _graph_tags(graph_path: str | None) -> dict[str, set[str]]:
    if not graph_path:
        return {}
    graph = load_json(graph_path)
    if not isinstance(graph.get("edges"), list):
        raise SystemExit(f"Career graph has no edges array: {Path(graph_path).expanduser()}")
    tags: dict[str, set[str]] = {}
    for edge in graph["edges"]:
        if not isinstance(edge, dict) or edge.get("type") != "tagged":
            continue
        source, target = str(edge.get("from") or ""), str(edge.get("to") or "")
        if source.startswith("achievement:") and target.startswith("tag:"):
            tags.setdefault(source.removeprefix("achievement:"), set()).add(target.removeprefix("tag:"))
    return tags


def _render_evidence(selected: list[Any], graph_tags: dict[str, set[str]], jd_source: str) -> str:
    lines = ["# Evidence map", "", f"- Job description: `{jd_source}`", f"- Selected achievements: {len(selected)}", ""]
    for item in selected:
        ach_id = str(getattr(item, "ach_id", ""))
        tags = sorted(graph_tags.get(ach_id, set(getattr(item, "tags", []) or [])))
        score = float(getattr(item, "score", 0.0))
        lines.extend([
            f"## {getattr(item, 'company', 'Unknown')} - {ach_id} (score {score:.3f})",
            f"- Tags: {', '.join(tags) if tags else 'none'}",
            f"- {str(getattr(item, 'text', '')).strip()}",
            "",
        ])
    return "\n".join(lines).rstrip() + "\n"


def build_career_materials(
    *,
    career_cli_path: str,
    career_bank_path: str,
    career_graph_path: str | None,
    role: dict[str, Any],
    role_path: Path,
    outdir: Path,
    write_docx: bool,
    top: int,
) -> tuple[Path, Path, Path | None]:
    cli = load_career_cli(career_cli_path)
    bank = load_json(career_bank_path)
    jd_text, jd_source = _load_job_description(role, role_path)
    if not jd_text.strip():
        raise SystemExit(f"Job description is empty: {jd_source}")
    selected = cli.tailor(bank, jd_text, top_n=top)
    if not selected:
        raise SystemExit("No career-bank achievements matched this job description; no resume was generated")
    graph_tags = _graph_tags(career_graph_path)
    company = str(role.get("company") or "Company")
    title = str(role.get("title") or "Role")
    safe_company = "".join(ch for ch in company if ch.isalnum() or ch in " -_").strip().replace(" ", "_")
    prefix = safe_company or "Company"
    resume_path = outdir / f"Resume_{prefix}.md"
    evidence_path = outdir / f"Evidence_{prefix}.md"
    resume_path.write_text(cli.render_resume_markdown(bank, selected, title=title), encoding="utf-8")
    evidence_path.write_text(_render_evidence(selected, graph_tags, jd_source), encoding="utf-8")
    docx_path: Path | None = None
    if write_docx:
        try:
            docx_path = outdir / f"Resume_{prefix}.docx"
            docx_path.write_bytes(cli.render_resume_docx_bytes(bank, selected, title=title))
        except ImportError as exc:
            raise SystemExit("python-docx is required for --docx output. Install it with: pip install python-docx") from exc
    return resume_path, evidence_path, docx_path


def main() -> None:
    ap = argparse.ArgumentParser(description="Tailor application materials from config or a local career bank.")
    ap.add_argument("--config", help="Path to config.json (required outside career mode)")
    ap.add_argument("--role", required=True, help="Path to role.json (company, title, highlights)")
    ap.add_argument("--out", default="./output", help="Output directory (default: ./output)")
    ap.add_argument("--career-bank", help="Path to a local career_bank.json; enables career mode")
    ap.add_argument("--career-graph", help="Optional local career_graph.json used for evidence tags")
    ap.add_argument("--career-cli", help="Path to the local career_cli.py used for ranking and rendering")
    ap.add_argument("--docx", action="store_true", help="Also write a DOCX resume in career mode")
    ap.add_argument("--top", type=int, default=12, help="Maximum career-bank achievements to select (default: 12)")
    args = ap.parse_args()

    if args.top <= 0:
        ap.error("--top must be positive")
    role_path = Path(args.role).expanduser()
    role = load_json(str(role_path))
    outdir = Path(args.out)
    outdir.mkdir(parents=True, exist_ok=True)

    if args.career_bank:
        if not args.career_cli:
            ap.error("--career-cli is required with --career-bank")
        resume, evidence, docx = build_career_materials(
            career_cli_path=args.career_cli,
            career_bank_path=args.career_bank,
            career_graph_path=args.career_graph,
            role=role,
            role_path=role_path,
            outdir=outdir,
            write_docx=args.docx,
            top=args.top,
        )
        print(f"Created: {resume}")
        print(f"Created: {evidence}")
        if docx:
            print(f"Created: {docx}")
        return

    if not args.config:
        ap.error("--config is required outside career mode")
    cfg = load_json(args.config)

    cover = build_cover_letter(cfg, role, outdir)
    resume = build_resume(cfg, outdir)
    print(f"Created: {cover}")
    print(f"Created: {resume}")


if __name__ == "__main__":
    main()
